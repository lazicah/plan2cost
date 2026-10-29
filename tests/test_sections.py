from openpyxl import load_workbook

from plan2cost.takeoff.bill import by_bill, numbered, write_pdf, write_xlsx
from plan2cost.takeoff.rough import FloorData, estimate
from plan2cost.takeoff.sections import bill_name, bill_of, group_of
from plan2cost.takeoff.walls import Opening
from tests.test_rough import RATES


def floors():
    ops = [Opening("W1200", 1.2, 1.2, 12, "window", True, True), Opening("D900", 0.9, 2.1, 10, "door", False, True)]
    return [FloorData("ground", 80.0, 80.0, 39, 10.4, 17.2, 25.0, ops, ["Kitchen", "Bath/wc"]), FloorData("first", 58.0, 48.0, 20, 10.4, 14.5, 17.0, [], ["Bedroom", "Bath/wc"])]


def items():
    return estimate(floors(), {"storey_height_m": 3.0, "roof_type": "hip"}, RATES["rates"])[0]


def test_the_bill_runs_substructure_then_superstructure_then_finishes_then_services():
    its = items()
    bills = [i.bill for i in its]
    assert bills == sorted(bills) and bills[0] == 2 and bills[-1] == 15
    assert list(dict.fromkeys(group_of(b) for b in bills)) == ["SUBSTRUCTURE", "SUPERSTRUCTURE", "INTERNAL FINISHES", "SERVICES"]
    assert [bill_name(b) for b in (2, 3, 4, 5, 15)] == ["Substructure", "Frame", "Upper floors", "Roof", "Electrical installations"]


def test_every_line_lands_in_a_real_bill_and_prefixed_lines_are_placed_by_rule():
    assert bill_of("opening_w1200_nr") == 9 and bill_of("opening_d900_nr") == 10 and bill_of("lintel_w1200_nr") == 7 and bill_of("ground_slab_m3") == 2 and bill_of("no_such_line") == 16
    assert 16 not in {i.bill for i in items()}  # nothing is left to fall through to external works


def test_the_excel_bill_has_group_and_bill_headings_subtotals_and_a_summary(tmp_path):
    rows = numbered([i.__dict__ for i in items()])
    write_xlsx(tmp_path / "b.xlsx", rows, True, "NGN", "Bill of quantities", "r1")
    ws = load_workbook(tmp_path / "b.xlsx")["Bill"]
    col = [r[0] for r in ws.iter_rows(min_row=5, values_only=True) if isinstance(r[0], str)]
    assert col[0] == "SUBSTRUCTURE" and col[1].startswith("BILL No. 02 — SUBSTRUCTURE") and any(c.startswith("BILL No. 15 — ELECTRICAL") for c in col)
    assert col.index("SUPERSTRUCTURE") > col.index("SUBSTRUCTURE") and col.index("SERVICES") > col.index("INTERNAL FINISHES")
    sm = load_workbook(tmp_path / "b.xlsx")["Summary"]
    labels = [r[1] for r in sm.iter_rows(values_only=True) if r[1]]
    assert labels[0].startswith("02 Substructure") and labels[-1] == "Total"


def test_the_pdf_bill_is_written_with_its_sections(tmp_path):
    rows = numbered([i.__dict__ for i in items()])
    write_pdf(tmp_path / "b.pdf", rows, True, "NGN", "Bill of quantities", "r1")
    assert (tmp_path / "b.pdf").read_bytes().startswith(b"%PDF") and len(by_bill(rows)) >= 10
