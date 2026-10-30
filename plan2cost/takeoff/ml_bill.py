"""The bill of materials and labour as a document (Excel and PDF), laid out like the author's SIMPLIFIED BOQ: each element with its material lines and labour lines and a subtotal,
inside the bills of the BoQ in the same order, with a summary at the end. Estimated lines show their low to high range."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEAD = ("Item", "Material / description", "Qty", "Unit", "Unit rate", "Total cost", "Total low", "Total high")
NOTE = ("Bill of materials and labour from floor plans. Quantities of materials and labour come from the author's own practice (cement bags per m3, planks per m2, labour rates) applied to the "
        "measured and estimated quantities of the BoQ; estimated lines show a low to high range. Rates are editable. Excludes preliminaries, overheads, profit, contingencies and VAT.")
SHADE = PatternFill("solid", fgColor="EFECE4")


def display_rows(bm: dict) -> list[dict]:
    """The bill flattened into rows: group, bill and element headings, lines, and subtotals."""
    rows, group = [], ""
    for sec in bm["sections"]:
        if sec["group"] != group:
            rows.append({"type": "group", "text": sec["group"]})
            group = sec["group"]
        rows.append({"type": "bill", "text": f"BILL No. {sec['bill']:02d} — {sec['name'].upper()}"})
        for sub in sec["subs"]:
            rows.append({"type": "sub", "text": sub["name"]})
            rows += [{"type": "line", **ln} for ln in sub["lines"]]
            rows.append({"type": "subtotal", "text": f"SUBTOTAL – {sub['name']}", **{k: sub[k] for k in ("amount", "amount_low", "amount_high")}})
        if sec["carried"]:
            rows.append({"type": "sub", "text": "Items carried at the BoQ rate (supply and fix, excavation and other lines with no recipe)"})
            rows += [{"type": "line", **ln} for ln in sec["carried"]]
        rows.append({"type": "billtotal", "text": f"SUBTOTAL – Bill No. {sec['bill']:02d} {sec['name']}", **{k: sec[k] for k in ("amount", "amount_low", "amount_high")}})
    return rows


def write_xlsx(path: str | Path, bm: dict, currency: str, title: str, run_id: str) -> None:
    """Write the bill of materials and labour as an Excel workbook, sectioned like the BoQ, with a low to high range on estimated lines."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Materials and labour"
    ws["A1"], ws["A1"].font = f"{title} ({currency})", Font(bold=True, size=13)
    ws["A2"] = NOTE
    ws.append([])
    ws.append(list(HEAD))
    for c in ws[4]:
        c.font = Font(bold=True)
    for r in display_rows(bm):
        t = r["type"]
        if t in ("group", "bill", "sub"):
            ws.append([r["text"] if t != "sub" else "", r["text"] if t == "sub" else ""])
            ws[f"A{ws.max_row}"].font = Font(bold=t != "sub", size=12 if t == "group" else 11)
            ws[f"B{ws.max_row}"].font = Font(bold=True, italic=True)
        elif t == "line":
            k = ws.max_row + 1
            ws.append([r["no"], r["description"], r["qty"], r["unit"], r["rate"], f"=C{k}*E{k}", r["amount_low"], r["amount_high"]])
        else:
            ws.append(["", r["text"], "", "", "", r["amount"], r["amount_low"], r["amount_high"]])
            for c in ws[ws.max_row]:
                c.font, c.fill = Font(bold=True), SHADE
    ws.append([])
    ws.append(["", "TOTAL – materials and labour", "", "", "", bm["total"], bm["total_low"], bm["total_high"]])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (9, 64, 10, 9, 13, 16, 16, 16)):
        ws.column_dimensions[col].width = w
    for r in range(5, ws.max_row + 1):
        for col in "EFGH":
            ws[f"{col}{r}"].number_format = "#,##0.00"
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    sm = wb.create_sheet("Summary")
    sm.append([f"{title}: summary ({currency})"])
    sm["A1"].font = Font(bold=True, size=13)
    sm.append([])
    sm.append(["", "Bill", "Amount", "Low", "High"])
    group = ""
    for sec in bm["sections"]:
        if sec["group"] != group:
            sm.append([sec["group"]])
            sm[f"A{sm.max_row}"].font = Font(bold=True)
            group = sec["group"]
        sm.append(["", f"{sec['bill']:02d} {sec['name']}", sec["amount"], sec["amount_low"], sec["amount_high"]])
    sm.append(["", "Total", bm["total"], bm["total_low"], bm["total_high"]])
    for c in sm[sm.max_row]:
        c.font = Font(bold=True)
    sm.column_dimensions["B"].width = 44
    for col in "CDE":
        sm.column_dimensions[col].width = 18
    for r in range(4, sm.max_row + 1):
        for col in "CDE":
            sm[f"{col}{r}"].number_format = "#,##0.00"
    info = wb.create_sheet("Info")
    info.append(["run", run_id])
    info.append(["currency", currency])
    info.append(["recipes", bm["source"]])
    wb.save(path)


def write_pdf(path: str | Path, bm: dict, currency: str, title: str, run_id: str) -> None:
    """Write the bill of materials and labour as a PDF."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = getSampleStyleSheet()
    cell = st["BodyText"].clone("cell", fontSize=8, leading=10)
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;")  # noqa: E731
    data, heads, subs = [list(HEAD[:6]) + ["Range"]], [], []
    for r in display_rows(bm):
        t = r["type"]
        if t in ("group", "bill", "sub"):
            data.append([Paragraph(f"<b>{esc(r['text'])}</b>" if t != "sub" else f"<i>{esc(r['text'])}</i>", cell)] + [""] * 6)
            heads.append((len(data) - 1, t))
        elif t == "line":
            data.append([r["no"], Paragraph(esc(r["description"]), cell), f"{r['qty']:,}", r["unit"], f"{r['rate']:,.2f}", f"{r['amount']:,.2f}",
                         f"{r['amount_low']:,.0f} to {r['amount_high']:,.0f}" if r["amount_low"] != r["amount_high"] else ""])
        else:
            data.append(["", Paragraph(f"<b>{esc(r['text'])}</b>", cell), "", "", "", f"{r['amount']:,.2f}", f"{r['amount_low']:,.0f} to {r['amount_high']:,.0f}" if r["amount_low"] != r["amount_high"] else ""])
            subs.append(len(data) - 1)
    data.append(["", Paragraph("<b>TOTAL – materials and labour</b>", cell), "", "", "", f"{bm['total']:,.2f}", f"{bm['total_low']:,.0f} to {bm['total_high']:,.0f}"])
    table = Table(data, colWidths=[40, 330, 52, 40, 64, 84, 130], repeatRows=1)
    style = [("FONTSIZE", (0, 0), (-1, -1), 8), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efece4")),
             ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b9b5aa")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (2, 0), (2, -1), "RIGHT"), ("ALIGN", (4, 0), (6, -1), "RIGHT")]
    style += [x for h, t in heads for x in (("SPAN", (0, h), (-1, h)), ("BACKGROUND", (0, h), (-1, h), colors.HexColor("#e2dfd7" if t != "sub" else "#f6f5f1")))]
    style += [("BACKGROUND", (0, s), (-1, s), colors.HexColor("#efece4")) for s in subs]
    table.setStyle(TableStyle(style))
    SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28, title=title).build(
        [Paragraph(f"<b>{esc(title)}</b> ({currency})", st["Title"]), Paragraph(NOTE, cell), Spacer(1, 8), table])
