"""The bill as a document. A priced bill goes to a client (rates include the contractor's overhead, profit and other costs); an unpriced bill
(quantities only) goes out for tender and the contractor fills in the rates. Preliminaries, overhead and profit live inside the rates, so the
bill has no separate section for them. Items are numbered so a returned bill can be matched line by line without guessing. Lines sit under the bill
headings of the author's My Duplex bill: substructure, superstructure, internal finishes, services, external works, each bill with its subtotal and a summary at the end."""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from plan2cost.takeoff.sections import bill_name, group_of

HEAD = ("No", "BESMM4", "Description", "Unit", "Qty", "Rate", "Amount")
NOTE = ("Price every item. Rates must include overhead, profit, preliminaries and all other costs: there is no separate item for them. "
        "Do not change quantities or descriptions; qualify in a covering letter instead.")
BUDGET_NOTE = ("Budget estimate from floor plans. Lines marked 'measured' come from the plans; lines marked 'estimated' come from coefficients and show a low to high range. "
               "Not a bill for tender: quantities must be measured from full drawings first.")
SHADE = PatternFill("solid", fgColor="EFECE4")


def numbered(items: list[dict]) -> list[dict]:
    """The BoQ lines with a stable item number (their order in the bill)."""
    return [{"no": str(n), "besmm": i.get("besmm", ""), "description": i["description"], "unit": i["unit"], "qty": i["qty"], "rate": i.get("rate"), "key": i["key"],
             "basis": i.get("basis", "measured"), "qty_low": i.get("qty_low"), "qty_high": i.get("qty_high"), "bill": i.get("bill") or 0} for n, i in enumerate(items, 1)]


def ranged(rows: list[dict]) -> bool:
    """A bill with estimated lines carries their range; it is a budget estimate, not a bill to send out for tender."""
    return any(r.get("basis") == "estimated" for r in rows)


def by_bill(rows: list[dict]) -> list[tuple[int, str, str, list[dict]]]:
    """(bill number, bill name, group, rows) in bill order; a bill with no lines is left out."""
    out = []
    for n in sorted({r["bill"] for r in rows if r.get("bill")}):
        out.append((n, bill_name(n), group_of(n), [r for r in rows if r.get("bill") == n]))
    unplaced = [r for r in rows if not r.get("bill")]
    return out + ([(0, "Other items", "OTHER", unplaced)] if unplaced else [])


def write_xlsx(path: str | Path, rows: list[dict], priced: bool, currency: str, title: str, run_id: str) -> None:
    """Write the bill as an Excel workbook: a sheet of items under their bill headings with subtotals, and a summary sheet. Priced bills carry rates and amounts (with a low to high range when lines are estimated); unpriced bills leave the rate blank and the amount a formula, for contractors to tender on."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bill"
    ws["A1"], ws["A1"].font = f"{title} ({'priced' if priced else 'unpriced, for tender'}; {currency})", Font(bold=True, size=13)
    budget = priced and ranged(rows)
    ws["A2"] = BUDGET_NOTE if budget else NOTE
    ws.append([])
    extra = ("Basis", "Qty low", "Qty high", "Amount low", "Amount high") if budget else ()
    ws.append(list(HEAD) + list(extra))
    for c in ws[4]:
        c.font = Font(bold=True)
    subtotals: list[tuple[str, str, int]] = []  # (group, bill name, row of its subtotal)
    last_group = ""
    for n, name, group, bill_rows in by_bill(rows):
        if group != last_group:
            ws.append([group])
            ws[f"A{ws.max_row}"].font = Font(bold=True, size=12)
            last_group = group
        ws.append([f"BILL No. {n:02d} — {name.upper()}" if n else name.upper()])
        ws[f"A{ws.max_row}"].font = Font(bold=True)
        first = ws.max_row + 1
        for r in bill_rows:
            k = ws.max_row + 1
            row = [int(r["no"]), r["besmm"], r["description"], r["unit"], r["qty"], r["rate"] if priced and r["rate"] else None, f"=E{k}*F{k}"]
            if extra:
                lo, hi = (r["qty_low"] if r.get("qty_low") is not None else r["qty"]), (r["qty_high"] if r.get("qty_high") is not None else r["qty"])
                row += [r.get("basis", "measured"), lo, hi, f"=I{k}*F{k}", f"=J{k}*F{k}"]
            ws.append(row)
            ws[f"C{k}"].alignment = Alignment(wrap_text=True, vertical="top")
        last = ws.max_row
        ws.append(["", "", f"Bill No. {n:02d} — {name} carried to summary" if n else f"{name} carried to summary", "", "", "", f"=SUM(G{first}:G{last})"]
                  + (["", "", "", f"=SUM(K{first}:K{last})", f"=SUM(L{first}:L{last})"] if extra else []))
        for c in ws[ws.max_row]:
            c.font, c.fill = Font(bold=True), SHADE
        subtotals.append((group, f"{n:02d} {name}" if n else name, ws.max_row))
        ws.append([])
    for col, w in zip("ABCDEFGHIJKL", (6, 16, 80, 7, 11, 14, 16, 11, 11, 11, 16, 16)):
        ws.column_dimensions[col].width = w
    for r in range(5, ws.max_row + 1):
        for col in "FGKL" if extra else "FG":
            ws[f"{col}{r}"].number_format = "#,##0.00"
    sm = wb.create_sheet("Summary")
    sm.append([f"{title}: summary ({currency})"])
    sm["A1"].font = Font(bold=True, size=13)
    sm.append([])
    group = ""
    total_rows = []
    for g, label, row in subtotals:
        if g != group:
            sm.append([g])
            sm[f"A{sm.max_row}"].font = Font(bold=True)
            group = g
        sm.append(["", label, f"=Bill!G{row}"] + ([f"=Bill!K{row}", f"=Bill!L{row}"] if extra else []))
        total_rows.append(sm.max_row)
    sm.append(["", "Total", "=" + "+".join(f"C{r}" for r in total_rows)] + (["=" + "+".join(f"D{r}" for r in total_rows), "=" + "+".join(f"E{r}" for r in total_rows)] if extra else []))
    for c in sm[sm.max_row]:
        c.font = Font(bold=True)
    sm.column_dimensions["B"].width = 44
    for col in "CDE":
        sm.column_dimensions[col].width = 18
    for r in range(3, sm.max_row + 1):
        for col in "CDE":
            sm[f"{col}{r}"].number_format = "#,##0.00"
    info = wb.create_sheet("Info")
    info.append(["run", run_id])
    info.append(["currency", currency])
    wb.save(path)


def read_xlsx_bid(path: str) -> dict[str, dict]:
    """A returned bill, by item number. Values are read as the contractor left them; nothing is corrected."""
    ws = load_workbook(path, data_only=True)["Bill"] if "Bill" in load_workbook(path, read_only=True).sheetnames else load_workbook(path, data_only=True).active
    head = next((i for i, row in enumerate(ws.iter_rows(values_only=True), 1) if row and row[0] == "No"), None)
    out: dict[str, dict] = {}
    if not head:
        return out
    for row in ws.iter_rows(min_row=head + 1, values_only=True):
        if not isinstance(row[0], (int, float)):
            continue
        qty, rate, amount = (float(v) if isinstance(v, (int, float)) else None for v in (row[4], row[5], row[6]))
        out[str(int(row[0]))] = {"description": str(row[2] or ""), "unit": str(row[3] or ""), "qty": qty, "rate": rate, "amount": amount}
    return out



def write_pdf(path: str | Path, rows: list[dict], priced: bool, currency: str, title: str, run_id: str) -> None:
    """The same bill as a PDF for reading and printing, in the same order with a subtotal per bill and a summary at the end. The unpriced PDF has empty rate and amount columns."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = getSampleStyleSheet()
    cell = st["BodyText"].clone("cell", fontSize=8, leading=10)
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28, title=title)
    rng = priced and ranged(rows)
    ncol = 9 if rng else 7
    data = [list(HEAD) + (["Basis", "Qty low to high"] if rng else [])]
    heads, subs, summary, total = [], [], [], 0.0
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;")  # noqa: E731
    group = ""
    for n, name, g, bill_rows in by_bill(rows):
        if g != group:
            data.append([Paragraph(f"<b>{g}</b>", cell)] + [""] * (ncol - 1))
            heads.append(len(data) - 1)
            group = g
        data.append([Paragraph(f"<b>BILL No. {n:02d} — {esc(name.upper())}</b>" if n else f"<b>{esc(name.upper())}</b>", cell)] + [""] * (ncol - 1))
        heads.append(len(data) - 1)
        bill_total = 0.0
        for r in bill_rows:
            amount = r["qty"] * r["rate"] if priced and r["rate"] else None
            bill_total += amount or 0
            data.append([r["no"], r["besmm"], Paragraph(esc(r["description"]), cell), r["unit"], f"{r['qty']:,.2f}", f"{r['rate']:,.2f}" if priced and r["rate"] else "", f"{amount:,.2f}" if amount else ""]
                        + ([r.get("basis", "measured"), f"{r['qty_low']:,.1f} to {r['qty_high']:,.1f}" if r.get("qty_low") is not None else ""] if rng else []))
        data.append(["", "", Paragraph(f"<b>Bill No. {n:02d} — {esc(name)} carried to summary</b>" if n else f"<b>{esc(name)} carried to summary</b>", cell), "", "", "", f"{bill_total:,.2f}" if priced else ""] + (["", ""] if rng else []))
        subs.append(len(data) - 1)
        summary.append((g, f"{n:02d} {name}" if n else name, bill_total))
        total += bill_total
    widths = [28, 62, 330, 34, 56, 64, 80] + ([52, 80] if rng else [])
    table = Table(data, colWidths=widths, repeatRows=1)
    style = [("FONTSIZE", (0, 0), (-1, -1), 8), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efece4")),
             ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b9b5aa")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (4, 0), (6, -1), "RIGHT")]
    style += [x for h in heads for x in (("SPAN", (0, h), (-1, h)), ("BACKGROUND", (0, h), (-1, h), colors.HexColor("#e2dfd7")))]
    style += [("BACKGROUND", (0, s), (-1, s), colors.HexColor("#efece4")) for s in subs]
    table.setStyle(TableStyle(style))
    sdata = [["", "Summary", f"Amount ({currency})"]]
    group = ""
    for g, label, amount in summary:
        if g != group:
            sdata.append([Paragraph(f"<b>{g}</b>", cell), "", ""])
            group = g
        sdata.append(["", label, f"{amount:,.2f}" if priced else ""])
    sdata.append(["", Paragraph("<b>Total</b>", cell), f"{total:,.2f}" if priced else ""])
    stable = Table(sdata, colWidths=[120, 300, 120])
    stable.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b9b5aa")), ("ALIGN", (2, 0), (2, -1), "RIGHT"),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efece4"))]))
    doc.build([Paragraph(f"<b>{title}</b> ({'priced' if priced else 'unpriced, for tender'}; {currency})", st["Title"]), Paragraph(BUDGET_NOTE if rng else NOTE, cell), Spacer(1, 8), table, Spacer(1, 14), stable])
